import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets internal cell margins in twips (1 pt = 20 twips)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_left_border_only(cell, color_hex="9E6C15", sz="36"):
    """Adds a thick accent border on the left side of a cell and removes other borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_table_borders(table, color_hex="D2A63C", sz="4"):
    """Sets thin subtle borders for the entire table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_house_rules_docx(output_path):
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("LULU AURELIAN ESTATE  •  GUEST HOUSE RULES & RESIDENTIAL MANUAL")
        hrun.font.name = "Georgia"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(158, 108, 21) # Amber gold

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Lulu Aurelian Furnished Apartments  •  Nyeri, Kenya  •  Contact Safaricom: 0112299384 / Airtel: 0756958531")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 115, 105)

    # Base Colors
    COLOR_PRIMARY = RGBColor(29, 25, 18)     # Deep Obsidian Charcoal #1D1912
    COLOR_GOLD = RGBColor(158, 108, 21)       # Amber Gold #9E6C15
    COLOR_GOLD_DARK = RGBColor(130, 85, 12)   # Dark Gold
    COLOR_MUTED = RGBColor(95, 90, 82)        # Warm Gray
    COLOR_WHITE = RGBColor(255, 255, 255)

    # -------------------------------------------------------------
    # COVER / HEADER SECTION
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    run_brand = title_p.add_run("LULU AURELIAN ESTATE")
    run_brand.font.name = "Georgia"
    run_brand.font.size = Pt(24)
    run_brand.font.bold = True
    run_brand.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("OFFICIAL HOUSE RULES & GUEST STAY GUIDELINES")
    run_sub.font.name = "Georgia"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_GOLD

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(16)
    run_meta = meta_p.add_run("Comprehensive Directory for All Residential Units: Skyview Hideaway (Penthouse), Cocoa Retreat (Apt 19), and Neema Haven (Apt SL-2)")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    run_meta.font.color.rgb = COLOR_MUTED

    # Welcome Callout Box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.8)
    cell = table.cell(0, 0)
    set_cell_background(cell, "FBF9F5") # Cream warm tint
    set_cell_left_border_only(cell, "9E6C15", "40")
    set_cell_margins(cell, top=160, bottom=160, left=220, right=180)

    box_p = cell.paragraphs[0]
    box_p.paragraph_format.space_before = Pt(0)
    box_p.paragraph_format.space_after = Pt(4)
    b_title = box_p.add_run("Hospitality Philosophy & Guest Commitment\n")
    b_title.font.name = "Georgia"
    b_title.font.size = Pt(11)
    b_title.font.bold = True
    b_title.font.color.rgb = COLOR_GOLD

    b_text = box_p.add_run(
        "At Lulu Aurelian Estate, we take pride in offering an atmosphere of effortless tranquility, refinement, "
        "and warmth. These house rules are designed to protect your serenity, preserve the pristine quality of our "
        "furnishings, and foster harmony with our neighbors. We kindly ask every guest to honor these guidelines during their stay."
    )
    b_text.font.name = "Calibri"
    b_text.font.size = Pt(10)
    b_text.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # QUICK COMPARISON SUMMARY TABLE
    # -------------------------------------------------------------
    sum_head = doc.add_paragraph()
    sum_head.paragraph_format.space_before = Pt(14)
    sum_head.paragraph_format.space_after = Pt(6)
    s_run = sum_head.add_run("Executive Overview: Units at a Glance")
    s_run.font.name = "Georgia"
    s_run.font.size = Pt(13)
    s_run.font.bold = True
    s_run.font.color.rgb = COLOR_PRIMARY

    tbl = doc.add_table(rows=4, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, "D2A63C", "6")

    col_widths = [Inches(1.5), Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.4)]
    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    headers = ["Unit Name", "Floor / Location", "Key Attributes", "Smoking Policy", "Quiet Hours"]
    for i, title in enumerate(headers):
        c = tbl.cell(0, i)
        set_cell_background(c, "1D1912")
        set_cell_margins(c, top=140, bottom=140, left=120, right=120)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Georgia"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = COLOR_WHITE

    units_data = [
        ("Skyview Hideaway", "Top Floor (Penthouse)", "Panoramic views, live plants, twin balconies", "Balcony only (ashtray provided)", "10:00 PM – 7:00 AM"),
        ("Cocoa Retreat", "First Floor (Apt 19)", "Warm cocoa hues, serene, private balcony", "Balcony only (ashtray provided)", "10:00 PM – 7:00 AM"),
        ("Neema Haven", "First Floor (Apt SL-2)", "Breezy sanctuary, tranquil rest, secure entry", "Balcony only (ashtray provided)", "10:00 PM – 7:00 AM"),
    ]

    for row_idx, row_data in enumerate(units_data, start=1):
        bg = "FFFFFF" if row_idx % 2 != 0 else "F8F6F0"
        for col_idx, text in enumerate(row_data):
            c = tbl.cell(row_idx, col_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=100, right=100)
            p = c.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.font.name = "Georgia"
                r.font.size = Pt(9)
                r.font.bold = True
                r.font.color.rgb = COLOR_GOLD
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_PRIMARY

    # Helper function for rendering unit house rules
    def add_unit_rules(unit_title, location, subtitle, welcome, gateway, highlights, sections, closing, contacts):
        doc.add_page_break()

        # Unit Header Title
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(10)
        hp.paragraph_format.space_after = Pt(2)
        hrun = hp.add_run(unit_title.upper())
        hrun.font.name = "Georgia"
        hrun.font.size = Pt(20)
        hrun.font.bold = True
        hrun.font.color.rgb = COLOR_PRIMARY

        lp = doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(0)
        lp.paragraph_format.space_after = Pt(4)
        lrun = lp.add_run(f"Location: {location}")
        lrun.font.name = "Georgia"
        lrun.font.size = Pt(11)
        lrun.font.bold = True
        lrun.font.color.rgb = COLOR_GOLD

        subp = doc.add_paragraph()
        subp.paragraph_format.space_before = Pt(0)
        subp.paragraph_format.space_after = Pt(10)
        sbrun = subp.add_run(subtitle)
        sbrun.font.name = "Calibri"
        sbrun.font.size = Pt(10.5)
        sbrun.font.italic = True
        sbrun.font.color.rgb = COLOR_MUTED

        # Welcome Box
        w_tbl = doc.add_table(rows=1, cols=1)
        w_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        w_tbl.columns[0].width = Inches(6.8)
        w_cell = w_tbl.cell(0, 0)
        set_cell_background(w_cell, "FBF9F5")
        set_cell_left_border_only(w_cell, "9E6C15", "36")
        set_cell_margins(w_cell, top=140, bottom=140, left=180, right=160)
        
        w_para = w_cell.paragraphs[0]
        w_para.paragraph_format.space_before = Pt(0)
        w_para.paragraph_format.space_after = Pt(4)
        w_r1 = w_para.add_run("A Warm Welcome to Our Guests\n")
        w_r1.font.name = "Georgia"
        w_r1.font.size = Pt(10.5)
        w_r1.font.bold = True
        w_r1.font.color.rgb = COLOR_GOLD

        w_r2 = w_para.add_run(welcome)
        w_r2.font.name = "Calibri"
        w_r2.font.size = Pt(9.5)
        w_r2.font.color.rgb = COLOR_PRIMARY

        if gateway:
            w_r3 = w_para.add_run("\n\n" + gateway)
            w_r3.font.name = "Calibri"
            w_r3.font.size = Pt(9.5)
            w_r3.font.italic = True
            w_r3.font.color.rgb = COLOR_PRIMARY

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Quick Highlights Table (2x2 grid)
        h_title_p = doc.add_paragraph()
        h_title_p.paragraph_format.space_before = Pt(8)
        h_title_p.paragraph_format.space_after = Pt(4)
        ht_run = h_title_p.add_run("Essential Highlights & Quick Guidelines")
        ht_run.font.name = "Georgia"
        ht_run.font.size = Pt(12)
        ht_run.font.bold = True
        ht_run.font.color.rgb = COLOR_PRIMARY

        hl_tbl = doc.add_table(rows=2, cols=2)
        hl_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        hl_tbl.columns[0].width = Inches(3.35)
        hl_tbl.columns[1].width = Inches(3.35)
        set_table_borders(hl_tbl, "E2DACF", "4")

        coords = [(0, 0), (0, 1), (1, 0), (1, 1)]
        for idx, (r_idx, c_idx) in enumerate(coords):
            if idx < len(highlights):
                hl = highlights[idx]
                c = hl_tbl.cell(r_idx, c_idx)
                set_cell_background(c, "FAF8F4")
                set_cell_margins(c, top=100, bottom=100, left=120, right=120)
                cp = c.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(2)
                r_t = cp.add_run(f"• {hl['title']}\n")
                r_t.font.name = "Georgia"
                r_t.font.size = Pt(9.5)
                r_t.font.bold = True
                r_t.font.color.rgb = COLOR_GOLD

                r_desc = cp.add_run(hl['text'])
                r_desc.font.name = "Calibri"
                r_desc.font.size = Pt(8.5)
                r_desc.font.color.rgb = COLOR_PRIMARY

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Full Detailed Rules Sections
        s_title_p = doc.add_paragraph()
        s_title_p.paragraph_format.space_before = Pt(8)
        s_title_p.paragraph_format.space_after = Pt(6)
        st_run = s_title_p.add_run("Detailed House Rules & Property Expectations")
        st_run.font.name = "Georgia"
        st_run.font.size = Pt(12)
        st_run.font.bold = True
        st_run.font.color.rgb = COLOR_PRIMARY

        for sec in sections:
            sec_p = doc.add_paragraph()
            sec_p.paragraph_format.space_before = Pt(6)
            sec_p.paragraph_format.space_after = Pt(2)
            sec_run = sec_p.add_run(sec['title'])
            sec_run.font.name = "Georgia"
            sec_run.font.size = Pt(10.5)
            sec_run.font.bold = True
            sec_run.font.color.rgb = COLOR_GOLD_DARK

            for item in sec['items']:
                ip = doc.add_paragraph()
                ip.paragraph_format.left_indent = Inches(0.25)
                ip.paragraph_format.space_before = Pt(0)
                ip.paragraph_format.space_after = Pt(2)
                irun_bullet = ip.add_run("✔  ")
                irun_bullet.font.name = "Segoe UI Symbol"
                irun_bullet.font.size = Pt(8.5)
                irun_bullet.font.color.rgb = COLOR_GOLD

                irun_text = ip.add_run(item)
                irun_text.font.name = "Calibri"
                irun_text.font.size = Pt(9.5)
                irun_text.font.color.rgb = COLOR_PRIMARY

        # Closing Note
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        c_tbl = doc.add_table(rows=1, cols=1)
        c_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_tbl.columns[0].width = Inches(6.8)
        c_cell = c_tbl.cell(0, 0)
        set_cell_background(c_cell, "1D1912")
        set_cell_margins(c_cell, top=140, bottom=140, left=180, right=180)

        cp = c_cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(3)
        cr_title = cp.add_run("A Warm Note of Gratitude\n")
        cr_title.font.name = "Georgia"
        cr_title.font.size = Pt(10)
        cr_title.font.bold = True
        cr_title.font.color.rgb = RGBColor(238, 205, 92) # Bright Gold

        cr_text = cp.add_run(closing + "\n\n")
        cr_text.font.name = "Calibri"
        cr_text.font.size = Pt(9)
        cr_text.font.color.rgb = COLOR_WHITE

        cr_contact = cp.add_run(f"Direct Concierge & Host Lines: Safaricom: {contacts['safaricom']}  |  Airtel: {contacts['airtel']}")
        cr_contact.font.name = "Calibri"
        cr_contact.font.size = Pt(9)
        cr_contact.font.bold = True
        cr_contact.font.color.rgb = RGBColor(238, 205, 92)

    # -------------------------------------------------------------
    # 1. SKYVIEW HIDEAWAY (TOP FLOOR PENTHOUSE)
    # -------------------------------------------------------------
    skyview_highlights = [
        {
            'title': 'Check-In & Out Schedule',
            'text': 'Check-in is from 2:00 PM • Check-out is at 10:00 AM so we can get the penthouse sparkling clean for our next arriving guests. (Need extra time? Just ask!)'
        },
        {
            'title': 'Fresh & Smoke-Free Indoors',
            'text': 'We keep indoor air fresh and clean for everyone. You are warmly welcome to enjoy the open mountain breeze on the balcony with the provided ashtrays.'
        },
        {
            'title': 'Peaceful Mountain Nights',
            'text': 'Evenings here are wonderfully quiet. We observe restful hours from 10:00 PM to 7:00 AM so everyone gets a great night’s sleep.'
        },
        {
            'title': 'Caring for the Penthouse & Towels',
            'text': 'Fluffy white towels and soft linens are ready for your comfort. A quick rinse of kitchen dishes before you head out helps our housekeeping crew immensely!'
        }
    ]

    skyview_sections = [
        {
            'title': '1. Make Yourself at Home & Protect Furnishings',
            'items': [
                'Please treat the apartment, cozy furnishings, and curated decor with the same love you’d treat your own home.',
                'Accidents happen—if anything spills or breaks, please let us know right away! Early communication helps us fix or replace items smoothly without any stress.',
                'Feel free to use the coasters and placemats on tables to keep the wood and stone surfaces looking pristine.',
                'All appliances, Smart TVs, high-speed Wi-Fi, and kitchen amenities are here for your enjoyment—please use them as intended.'
            ]
        },
        {
            'title': '2. Keeping Things Fresh, Towel Etiquette & Kitchen Care',
            'items': [
                'We love keeping the living space and open kitchen welcoming for you—a quick tidy-up of personal items as you go keeps the space feeling open and relaxing.',
                'Before checking out, a quick rinse or wash of your used plates, cups, and cookware is deeply appreciated by our cleaning crew.',
                'Trash bins are placed conveniently throughout the home—please pop waste inside and wipe up any quick spills.',
                'Our plush white towels are provided exclusively for your personal drying and bathing comfort. To help us keep them fresh and stain-free, please do not use them to wipe down shoes, luggage, makeup, or floor spills.',
                'Leaving the home reasonably tidy helps us turn it around promptly and beautifully for the next arriving guest.'
            ]
        },
        {
            'title': '3. Peaceful Nights & Good Neighbors',
            'items': [
                'Sound carries softly in the fresh mountain air. We kindly ask for quiet, restful hours between 10:00 PM and 7:00 AM.',
                'Our neighbors in the building are lovely families and professionals—thank you for helping us maintain a peaceful, respectful neighborhood atmosphere.'
            ]
        },
        {
            'title': '4. Smoke-Free Sanctuary',
            'items': [
                'To keep the apartment smelling fresh, crisp, and allergen-free for all travelers, indoor smoking and vaping are strictly prohibited.',
                'You are warmly welcome to step out onto the private balcony to smoke—please use the provided ashtrays and keep the balcony door closed so smoke doesn’t drift inside.'
            ]
        },
        {
            'title': '5. Visitors, Gatherings & Reservation Policy',
            'items': [
                'Your booking covers the guests listed on your reservation. If you’d like to host a friend for coffee or lunch, just give us a quick heads-up beforehand!',
                'Our home is intended as a calm sanctuary, so large parties, loud celebrations, and unauthorized overnight visitors are strictly prohibited.'
            ]
        },
        {
            'title': '6. Furry Friends & Pets',
            'items': [
                'We love pets! Just be sure to check in with management in advance so we can confirm if your four-legged companion can be accommodated.',
                'Please make sure pets stay off the beds, don’t disturb our neighbors, and that any messes are cleaned up promptly.'
            ]
        },
        {
            'title': '7. Your Safety, Security & Peace of Mind',
            'items': [
                'For your comfort and peace of mind, please remember to lock the front door and latch the windows whenever you head out exploring.',
                'Smoke detectors and safety devices are installed to protect you—please keep them clear and undisturbed.',
                'Take a quick moment to familiarize yourself with the building stairs and emergency exits when you settle in.'
            ]
        },
        {
            'title': '8. Stress-Free Checkout Procedure',
            'items': [
                'Check-out is at 10:00 AM so we can get the space sparkling clean for our next guests. If you need a bit of extra time, just let us know and we’ll do our best to accommodate you!',
                'Please place your used towels together in the bathroom basket or laundry hamper.',
                'Give the lights, electronics, and heating or AC a quick switch-off before locking the door behind you.'
            ]
        },
        {
            'title': '9. Our Green Friends (Live Indoor Plants)',
            'items': [
                'You’ll notice several live plants around the apartment—they bring fresh oxygen, nature, and positive energy into the penthouse!',
                'Our team takes care of their regular watering and nourishment, so there is no need to water or move them. Please ensure little hands don’t tug on the leaves.',
                'Your gentle appreciation keeps them happy and thriving!'
            ]
        },
        {
            'title': '10. Penthouse Balcony, High Views & Laundry Machine',
            'items': [
                'The top-floor balcony offers panoramic views of Nyeri’s rolling hills—the ideal spot for morning coffee or evening sunsets.',
                'Please always keep a watchful eye on young children whenever they enjoy the balcony.',
                'For your safety, please avoid leaning over the railings or climbing on balcony furniture.',
                'Everything in the apartment is here for your enjoyment! If you’d like to use the washing machine located on the second balcony, just send us a quick text and our team will gladly get you set up.'
            ]
        }
    ]

    add_unit_rules(
        unit_title="Skyview Hideaway",
        location="Top Floor (Penthouse)",
        subtitle="Welcome to Lulu Aurelian Apartment Skyview Hideaway (Top Floor)!",
        welcome="We’re so happy to welcome you! This mountain retreat was prepared with care and comfort in mind so you can unwind and feel right at home. Here are a few friendly guidelines to ensure your stay—and everyone else’s—is smooth, safe, and truly relaxing.",
        gateway="Welcome to Lulu Aurelian Apartments - Skyview Gateway:\nWe’re honored to share our top-floor apartment with you. This space has been prepared with care, thoughtfulness, and a genuine love for comfort. To help us maintain its beauty for everyone who stays, we kindly invite you to review our guidelines below.",
        highlights=skyview_highlights,
        sections=skyview_sections,
        closing="Thank you so much for being a thoughtful and considerate guest. Your care helps us keep this home a joyful retreat for everyone. We hope your time here is peaceful, memorable, and filled with warm moments!",
        contacts={'safaricom': '0112299384', 'airtel': '0756958531'}
    )

    # -------------------------------------------------------------
    # 2. COCOA RETREAT (FIRST FLOOR / APARTMENT 19)
    # -------------------------------------------------------------
    cocoa_highlights = [
        {
            'title': 'Check-In & Out Schedule',
            'text': 'Check-in is from 2:00 PM • Check-out is at 10:00 AM so we can get the space sparkling clean for our next guests. (Need extra time? Just ask!)'
        },
        {
            'title': 'Fresh & Smoke-Free Indoors',
            'text': 'We keep indoor air clean and inviting. Outdoor balcony space with an ashtray is ready whenever you need fresh air.'
        },
        {
            'title': 'Peaceful Living Community',
            'text': 'Our residential community observes quiet hours between 10:00 PM and 7:00 AM so you can enjoy deep, uninterrupted rest.'
        },
        {
            'title': 'Comfort & Towel Care',
            'text': 'Fresh, fluffy towels are ready for you. A quick rinse of your dishes before you check out makes our cleaning team smile!'
        }
    ]

    cocoa_sections = [
        {
            'title': '1. Caring for Your Home Away from Home',
            'items': [
                'Please enjoy and treat the plush cocoa-toned furnishings, artwork, and decor with gentle care.',
                'Spills and accidents happen—if anything gets damaged, please let us know right away! Early communication makes it simple and easy for us to take care of it.',
                'Kindly use the provided coasters and mats to protect our lovely tabletops and counters.',
                'Kitchen appliances and entertainment features are all yours to enjoy—please use them as intended.'
            ]
        },
        {
            'title': '2. Cleanliness, Kitchen Care & White Towel Usage',
            'items': [
                'A quick tidy-up of kitchen counters and living spaces as you go keeps the retreat feeling luxurious and open.',
                'Please give used dishes and cookware a quick wash or rinse before your departure.',
                'Trash bins are ready for your use—please toss waste inside and wipe up spills while they’re fresh.',
                'Our white towels are provided exclusively for your personal drying and bathing comfort. To keep them fresh and stain-free for every guest, please avoid using them to wipe down shoes, luggage, or spills.',
                'Leaving the apartment reasonably tidy helps our housekeeping team prepare the space smoothly for the next traveler.'
            ]
        },
        {
            'title': '3. Peaceful Nights & Good Neighbors',
            'items': [
                'Please keep music and conversation to a respectful volume, particularly during quiet hours from 10:00 PM to 7:00 AM.',
                'Our neighbors in the building appreciate the tranquil atmosphere—thank you for being mindful of those around you.'
            ]
        },
        {
            'title': '4. Smoke-Free Sanctuary',
            'items': [
                'To keep the retreat smelling sweet and fresh for everyone, indoor smoking and vaping are not permitted.',
                'You are welcome to smoke outside on the balcony—please use the provided ashtrays and close the sliding door so smoke stays outside.'
            ]
        },
        {
            'title': '5. Visitors & Friends Policy',
            'items': [
                'Your stay covers the guests on your booking. If you plan to invite a friend over for tea or dinner, please let us know beforehand!',
                'Cocoa Retreat is a tranquil space, so we do not host parties, large gatherings, or unregistered overnight guests.'
            ]
        },
        {
            'title': '6. Your Safety & Security',
            'items': [
                'Please lock the entry door and latch windows whenever you leave the retreat.',
                'Safety devices and fire equipment are installed for your security—please keep them undisturbed.',
                'Your lockbox passcode is private for your booking—please keep it secure.'
            ]
        },
        {
            'title': '7. Stress-Free Checkout',
            'items': [
                'Check-out is at 10:00 AM so we can get the space sparkling clean for our next guests. If you need a bit of extra time, just let us know!',
                'Please place your used towels in the bathroom laundry basket.',
                'Before you close the door, please ensure kitchen appliances, lights, and electronics are turned off.'
            ]
        }
    ]

    add_unit_rules(
        unit_title="Cocoa Retreat",
        location="First Floor (Apartment 19)",
        subtitle="Welcome to Lulu Aurelian Furnished Apartments - Cocoa Retreat!",
        welcome="We’re so delighted to have you stay with us at Cocoa Retreat. Designed for serene, sophisticated living, we invite you to review these friendly guidelines so your stay is completely relaxing, seamless, and comfortable.",
        gateway="",
        highlights=cocoa_highlights,
        sections=cocoa_sections,
        closing="Thank you for choosing Cocoa Retreat and treating this home with genuine love and care. We hope your stay is filled with comfort, calm, and memorable moments!",
        contacts={'safaricom': '0112299384', 'airtel': '0756958531'}
    )

    # -------------------------------------------------------------
    # 3. NEEMA HAVEN (FIRST FLOOR / APARTMENT SL-2)
    # -------------------------------------------------------------
    neema_highlights = [
        {
            'title': 'Check-In & Out Schedule',
            'text': 'Check-in is from 2:00 PM • Check-out is at 10:00 AM so we can get the space sparkling clean for our next guests. (Need extra time? Just ask!)'
        },
        {
            'title': 'Fresh & Smoke-Free Indoors',
            'text': 'We maintain a 100% smoke-free home indoors. You are welcome to smoke on the balcony with the provided ashtrays.'
        },
        {
            'title': 'Peaceful Haven Living',
            'text': 'Enjoy quiet, restful hours between 10:00 PM and 7:00 AM for a deeply refreshing night of sleep.'
        },
        {
            'title': 'Balcony Care & High Winds',
            'text': 'Please keep the balcony door securely latched during windy afternoon breezes.'
        }
    ]

    neema_sections = [
        {
            'title': '1. Make Yourself at Home',
            'items': [
                'Please treat the bedding, furnishings, and home amenities with kindness and care.',
                'If an accident or breakage occurs, just drop our team a quick text—early communication helps us resolve things without any worry!',
                'All appliances and comforts are here for you to enjoy—please use them as intended.'
            ]
        },
        {
            'title': '2. Keeping Things Fresh & Clean',
            'items': [
                'Tidying up after yourself in the kitchen and common areas keeps the haven feeling light and spacious.',
                'Please wash or rinse used cookware and dishes before you head out.',
                'Trash bins are ready for you—please dispose of waste inside and wipe up spills promptly.',
                'Our fluffy towels are provided exclusively for your personal drying and bathing comfort. Please ensure they are not used to clean shoes, luggage, or floor spills.'
            ]
        },
        {
            'title': '3. Peaceful Nights & Good Neighbors',
            'items': [
                'Please keep volume and music to a gentle level, especially during quiet hours from 10:00 PM to 7:00 AM.',
                'Thank you for being thoughtful of our neighbors throughout the building.'
            ]
        },
        {
            'title': '4. Smoke-Free Sanctuary',
            'items': [
                'To keep indoor air clean and fresh for every guest, smoking and vaping are not permitted inside.',
                'You are welcome to step onto the balcony to smoke—please use the provided ashtrays and close the door behind you.'
            ]
        },
        {
            'title': '5. Visitors & Friends',
            'items': [
                'Your stay is reserved for the guests listed on your booking. If you’d like to invite a guest over for a chat, please let us know beforehand.',
                'To preserve the peaceful haven for everyone, we do not host parties or unregistered overnight visitors.'
            ]
        },
        {
            'title': '6. Your Safety & Security (Balcony Wind Care)',
            'items': [
                'Always lock the front door and secure windows whenever you step out.',
                'When mountain winds pick up, please ensure the balcony door is firmly latched to avoid slam damage.',
                'Keep your lockbox passcode private for your own peace of mind.'
            ]
        },
        {
            'title': '7. Stress-Free Checkout',
            'items': [
                'Check-out is at 10:00 AM so we can get the space sparkling clean for our next guests. If you need a bit of extra time, just let us know!',
                'Please leave used towels in the designated bathroom basket.',
                'A quick turn-off of the lights and appliances before locking up is deeply appreciated.'
            ]
        }
    ]

    add_unit_rules(
        unit_title="Neema Haven",
        location="First Floor (Apartment SL-2)",
        subtitle="Welcome to Lulu Aurelian Furnished Apartments - Neema Haven!",
        welcome="A warm welcome to Neema Haven! Designed as a restful sanctuary, we invite you to review these friendly guidelines so your stay is wonderfully peaceful, safe, and rejuvenating.",
        gateway="",
        highlights=neema_highlights,
        sections=neema_sections,
        closing="Thank you for choosing Neema Haven and being such a thoughtful, wonderful guest. We hope your stay leaves you refreshed, inspired, and deeply rested!",
        contacts={'safaricom': '0112299384', 'airtel': '0756958531'}
    )

    # -------------------------------------------------------------
    # 4. GENERAL ESTATE-WIDE RULES & STANDARD POLICIES
    # -------------------------------------------------------------
    doc.add_page_break()

    ep_p = doc.add_paragraph()
    ep_p.paragraph_format.space_before = Pt(10)
    ep_p.paragraph_format.space_after = Pt(2)
    ep_run = ep_p.add_run("GENERAL ESTATE-WIDE POLICIES & PROCEDURES")
    ep_run.font.name = "Georgia"
    ep_run.font.size = Pt(20)
    ep_run.font.bold = True
    ep_run.font.color.rgb = COLOR_PRIMARY

    es_p = doc.add_paragraph()
    es_p.paragraph_format.space_before = Pt(0)
    es_p.paragraph_format.space_after = Pt(12)
    es_run = es_p.add_run("Applicable to all guests, visitors, and residents across Lulu Aurelian Estate")
    es_run.font.name = "Georgia"
    es_run.font.size = Pt(11)
    es_run.font.bold = True
    es_run.font.color.rgb = COLOR_GOLD

    general_sections = [
        {
            'title': 'A. Smart Lockbox Access & Key Security',
            'items': [
                'Each apartment is fitted with a secure key lockbox or digital lock. The secret 4-digit code is generated uniquely for your reservation and sent via SMS/WhatsApp on arrival day.',
                'Do not share your passcode with outside third parties or unauthorized visitors.',
                'When leaving the premises, ensure the physical keys remain inside the lockbox or are safely retained in your personal possession. Replacement of lost keys and lock changes incur a fee of KES 2,500.'
            ]
        },
        {
            'title': 'B. High-Speed Wi-Fi & Internet Acceptable Use',
            'items': [
                'Complimentary high-speed fiber internet is provided in each suite (Dedicated 5G/2.4G networks).',
                'Network credentials are displayed on the router card and in your digital welcome portal.',
                'Guests may not use the estate Wi-Fi for unlawful activities, commercial server hosting, or torrenting copyrighted material.'
            ]
        },
        {
            'title': 'C. Vehicle Parking & Traffic Etiquette',
            'items': [
                'Designated perimeter secure parking is available for guests within the gated compound with 24/7 security surveillance.',
                'Please park strictly inside marked parking bays and avoid blocking access gates, stairwells, or neighboring vehicle bays.',
                'The maximum estate speed limit is 10 km/h to protect strolling guests and children.'
            ]
        },
        {
            'title': 'D. Strict Prohibition of Large Parties & Unauthorized Commercial Events',
            'items': [
                'Lulu Aurelian Estate is a dedicated retreat and sanctuary for peaceful relaxation and business travelers.',
                'Commercial filming, photo shoots with large crews, stag/bachelor parties, and ticketed events are prohibited unless explicitly approved in writing with estate management.',
                'Violations of noise limits or party bans may result in immediate revocation of access without refund.'
            ]
        },
        {
            'title': 'E. Waste Disposal & Environmental Conservation',
            'items': [
                'Please dispose of household rubbish in designated tied refuse bags before placing them in the estate waste collection area.',
                'Conserve energy by switching off hot water switches, ceiling fans, and room lighting whenever you are stepping out of the apartment.'
            ]
        },
        {
            'title': 'F. Damages, Broken Items & Security Deposits',
            'items': [
                'Guests are held responsible for severe damage to linens, structural fixtures, electronics, or broken décor caused by negligence or misuse.',
                'Any unintentional breakages reported immediately to our management will be handled transparently at replacement cost with zero penalty markups.'
            ]
        }
    ]

    for gsec in general_sections:
        gp = doc.add_paragraph()
        gp.paragraph_format.space_before = Pt(7)
        gp.paragraph_format.space_after = Pt(2)
        grun = gp.add_run(gsec['title'])
        grun.font.name = "Georgia"
        grun.font.size = Pt(11)
        grun.font.bold = True
        grun.font.color.rgb = COLOR_GOLD_DARK

        for item in gsec['items']:
            ip = doc.add_paragraph()
            ip.paragraph_format.left_indent = Inches(0.25)
            ip.paragraph_format.space_before = Pt(0)
            ip.paragraph_format.space_after = Pt(2)
            irun_bullet = ip.add_run("✔  ")
            irun_bullet.font.name = "Segoe UI Symbol"
            irun_bullet.font.size = Pt(8.5)
            irun_bullet.font.color.rgb = COLOR_GOLD

            irun_text = ip.add_run(item)
            irun_text.font.name = "Calibri"
            irun_text.font.size = Pt(9.5)
            irun_text.font.color.rgb = COLOR_PRIMARY

    # Emergency & Contacts Box
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    e_tbl = doc.add_table(rows=1, cols=1)
    e_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    e_tbl.columns[0].width = Inches(6.8)
    e_cell = e_tbl.cell(0, 0)
    set_cell_background(e_cell, "FBF9F5")
    set_cell_left_border_only(e_cell, "9E6C15", "40")
    set_cell_margins(e_cell, top=140, bottom=140, left=180, right=160)

    epara = e_cell.paragraphs[0]
    epara.paragraph_format.space_before = Pt(0)
    epara.paragraph_format.space_after = Pt(2)
    e_r1 = epara.add_run("24/7 Estate Management & Emergency Contacts\n")
    e_r1.font.name = "Georgia"
    e_r1.font.size = Pt(11)
    e_r1.font.bold = True
    e_r1.font.color.rgb = COLOR_GOLD

    e_r2 = epara.add_run(
        "For immediate assistance, early check-in inquiries, extra linens, or housekeeping requests:\n"
        "• Host Concierge (Safaricom): +254 112 299 384\n"
        "• Host Concierge (Airtel): +254 756 958 531\n"
        "• Security Gate & Building Management: On-Duty 24/7 at Main Compound Gate\n"
        "• Emergency Services (Police / Fire / Medical): 999 / 112\n\n"
        "Thank you for making Lulu Aurelian Estate your home in Nyeri. We wish you an exceptional stay!"
    )
    e_r2.font.name = "Calibri"
    e_r2.font.size = Pt(9.5)
    e_r2.font.color.rgb = COLOR_PRIMARY

    # Save document
    doc.save(output_path)
    print(f"Successfully generated house rules Word document at: {output_path}")

if __name__ == "__main__":
    import os
    target_path = r"c:\wamp64\www\LULU AURELIAN ESTATE\Lulu_Aurelian_Estate_House_Rules.docx"
    build_house_rules_docx(target_path)
